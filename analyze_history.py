# -*- coding: utf-8 -*-
"""
历史数据分析脚本（增强版 + 量化打分引擎）
功能：读取 history_data(个股) 和 sector_data(板块) 里的数据，计算：
      基础表现：5/20/60日涨跌幅、量能变化
      趋势判断：均线排列、MACD金叉死叉、相对沪深300强弱
      建仓参考：RSI(14)、乖离率
      风险参考：ATR止损参考幅度
      【新增】量化打分引擎：参考常见的多因子选股框架，把"趋势/量能/板块"三大类
      信号拆解成具体加分项，汇总成一个 0-100 的综合评分(FinalScore)，并单独列出
      风险扣分项。这是该框架的简化落地版本——没有接入"大盘环境"和"集合竞价"两个
      模块（分别需要全市场涨跌家数数据和盘中竞价数据），所以分数上限并不是完整的
      100分，而是按已实现的模块重新换算到0-100，方便直观比较，但不能和以后补全
      全部模块的分数直接对比。

重要说明（务必先看）：
      本脚本所有指标和评分都基于公开、通用的技术分析公式和规则计算，是对历史价格
      规律的客观统计描述，不构成任何投资建议，也不保证未来表现。评分越高只代表
      "当前多个技术信号同时偏多"，不是"上涨概率"，两者是不同的概念——要得到真正
      的历史上涨概率，需要对评分分桶做历史回测校准，这是后续可以再做的一步。
      是否建仓、仓位多少、止损设在哪，最终都需要你自己判断和承担风险。脚本作者/
      生成者不是持牌投资顾问。

用法：
    1. 先运行 fetch_history.py 更新个股和沪深300数据
    2. 建议也运行一次 fetch_sector.py 获取板块归属和板块历史数据(不常变，一两周跑一次即可)
       如果没有运行过 fetch_sector.py，板块相关的评分项会自动跳过，不影响其他部分
    3. 运行本脚本：
       python analyze_history.py
    4. 生成 分析报表.xlsx 和 选股数据.json，双击打开Excel或用dashboard.html查看即可
"""

import pandas as pd
import numpy as np
import os
import glob
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

HISTORY_DIR = "history_data"
SECTOR_DIR = "sector_data"
OUTPUT_FILE = "分析报表.xlsx"
JSON_FILE = "选股数据.json"
BENCHMARK_CODE = "000300"


def load_all_files() -> dict:
    """读取 history_data 文件夹里所有csv，返回 {股票代码: DataFrame}（含基准指数）"""
    result = {}
    files = glob.glob(os.path.join(HISTORY_DIR, "*.csv"))
    for f in files:
        if "_合并历史数据" in f:
            continue
        df = pd.read_csv(f, encoding="utf-8-sig")
        if df.empty or "股票代码" not in df.columns:
            continue
        df["日期"] = pd.to_datetime(df["日期"])
        df = df.sort_values("日期").reset_index(drop=True)
        code = str(df["股票代码"].iloc[0]).zfill(6)
        result[code] = df
    return result


def load_sector_data():
    """
    读取 sector_data 文件夹（如果存在）
    返回：(股票代码->行业代码/名称 的映射dict, 行业代码->历史数据DataFrame 的dict)
    如果没跑过 fetch_sector.py，两个都返回空dict，不影响其他指标计算
    """
    stock_to_sector = {}
    sector_histories = {}

    map_path = os.path.join(SECTOR_DIR, "stock_industry_map.csv")
    if not os.path.exists(map_path):
        return stock_to_sector, sector_histories

    mapping = pd.read_csv(map_path, encoding="utf-8-sig", dtype={"股票代码": str})
    for _, row in mapping.iterrows():
        code = str(row["股票代码"]).zfill(6)
        stock_to_sector[code] = {
            "行业代码": str(row["行业代码"]),
            "行业名称": row.get("行业名称", row["行业代码"]),
        }

    for industry_code in mapping["行业代码"].astype(str).unique():
        hist_path = os.path.join(SECTOR_DIR, f"{industry_code}.csv")
        if os.path.exists(hist_path):
            df = pd.read_csv(hist_path, encoding="utf-8-sig")
            df["日期"] = pd.to_datetime(df["日期"])
            df = df.sort_values("日期").reset_index(drop=True)
            sector_histories[industry_code] = df

    return stock_to_sector, sector_histories


# ---------------- 技术指标计算函数 ----------------

def calc_ma(df, window):
    return df["收盘"].rolling(window).mean()


def calc_macd(close: pd.Series):
    """返回 DIF, DEA, MACD柱"""
    ema12 = close.ewm(span=12, adjust=False).mean()
    ema26 = close.ewm(span=26, adjust=False).mean()
    dif = ema12 - ema26
    dea = dif.ewm(span=9, adjust=False).mean()
    hist = (dif - dea) * 2
    return dif, dea, hist


def calc_rsi(close: pd.Series, period: int = 14):
    delta = close.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.rolling(period).mean()
    avg_loss = loss.rolling(period).mean()
    rs = avg_gain / avg_loss.replace(0, np.nan)
    rsi = 100 - (100 / (1 + rs))
    return rsi


def calc_atr(df: pd.DataFrame, period: int = 14):
    prev_close = df["收盘"].shift(1)
    tr1 = df["最高"] - df["最低"]
    tr2 = (df["最高"] - prev_close).abs()
    tr3 = (df["最低"] - prev_close).abs()
    tr = pd.concat([tr1, tr2, tr3], axis=1).max(axis=1)
    return tr.rolling(period).mean()


def calc_obv(df: pd.DataFrame):
    """能量潮指标：上涨日累加成交量，下跌日累减成交量"""
    direction = np.sign(df["收盘"].diff().fillna(0))
    return (direction * df["成交量"]).cumsum()


def percentile_rank(series: pd.Series, value: float) -> float:
    """计算 value 在 series 里的历史分位数(0-100)"""
    valid = series.dropna()
    if len(valid) == 0 or pd.isna(value):
        return None
    return round((valid < value).sum() / len(valid) * 100, 1)


# ---------------- 单只股票指标汇总 ----------------

def calc_metrics(df: pd.DataFrame, benchmark_df: pd.DataFrame,
                  sector_info: dict = None, sector_hist: pd.DataFrame = None) -> dict:
    if len(df) < 30:
        return None  # 数据不足30天，均线/MACD/RSI等指标意义不大，跳过

    latest = df.iloc[-1]
    name = df["股票名称"].iloc[-1]
    code = str(df["股票代码"].iloc[-1]).zfill(6)

    def pct_change(n):
        if len(df) <= n:
            return None
        old_close = df["收盘"].iloc[-(n + 1)]
        new_close = df["收盘"].iloc[-1]
        if not old_close or pd.isna(old_close):
            return None
        return round((new_close - old_close) / old_close * 100, 2)

    chg_5 = pct_change(5)
    chg_20 = pct_change(20)
    chg_60 = pct_change(60)

    # 量能变化
    vol_ratio = None
    if "成交量" in df.columns and len(df) >= 25:
        recent_5_vol = df["成交量"].iloc[-5:].mean()
        prior_20_vol = df["成交量"].iloc[-25:-5].mean()
        if prior_20_vol and prior_20_vol > 0:
            vol_ratio = round(recent_5_vol / prior_20_vol, 2)

    # 均线 & 排列判断
    ma5 = calc_ma(df, 5).iloc[-1]
    ma20 = calc_ma(df, 20).iloc[-1]
    ma60 = calc_ma(df, 60).iloc[-1] if len(df) >= 60 else None

    if ma60 is not None and not pd.isna(ma60):
        if ma5 > ma20 > ma60:
            ma_trend = "多头排列"
        elif ma5 < ma20 < ma60:
            ma_trend = "空头排列"
        else:
            ma_trend = "震荡交织"
    else:
        ma_trend = "数据不足(需60日+)"

    above_ma20 = "是" if latest["收盘"] > ma20 else "否"

    # 乖离率（距20日均线偏离度）
    bias20 = round((latest["收盘"] - ma20) / ma20 * 100, 2) if ma20 else None

    # MACD
    dif, dea, hist = calc_macd(df["收盘"])
    macd_status = "数据不足"
    if len(df) >= 35:
        if hist.iloc[-2] < 0 and hist.iloc[-1] > 0:
            macd_status = "金叉(转多)"
        elif hist.iloc[-2] > 0 and hist.iloc[-1] < 0:
            macd_status = "死叉(转空)"
        elif dif.iloc[-1] > dea.iloc[-1]:
            macd_status = "多头持续"
        else:
            macd_status = "空头持续"

    # RSI
    rsi = calc_rsi(df["收盘"]).iloc[-1]
    rsi = round(rsi, 1) if not pd.isna(rsi) else None
    if rsi is not None:
        if rsi >= 70:
            rsi_flag = "超买区间"
        elif rsi <= 30:
            rsi_flag = "超卖区间"
        else:
            rsi_flag = "中性"
    else:
        rsi_flag = "数据不足"

    # ATR 止损参考幅度（2倍ATR法，是常见的技术止损参考方法之一）
    atr = calc_atr(df).iloc[-1]
    stop_loss_pct = None
    stop_loss_price = None
    if atr and not pd.isna(atr) and latest["收盘"] > 0:
        stop_loss_pct = round((2 * atr) / latest["收盘"] * 100, 2)
        stop_loss_price = round(latest["收盘"] - 2 * atr, 2)

    # 60日区间高低点，给网页版画"价格区间位置条"用
    lookback60 = min(60, len(df))
    high_60 = round(df["最高"].iloc[-lookback60:].max(), 2)
    low_60 = round(df["最低"].iloc[-lookback60:].min(), 2)

    # 相对强弱（跟沪深300比，20日区间涨跌幅之差）
    rel_strength = None
    if benchmark_df is not None and len(benchmark_df) > 20 and chg_20 is not None:
        bench_old = benchmark_df["收盘"].iloc[-21]
        bench_new = benchmark_df["收盘"].iloc[-1]
        if bench_old and bench_old > 0:
            bench_chg_20 = (bench_new - bench_old) / bench_old * 100
            rel_strength = round(chg_20 - bench_chg_20, 2)

    # 综合趋势评分（0-5分，越高说明多个信号同时偏多，仅供参考，不是买卖信号）
    score = 0
    if ma_trend == "多头排列":
        score += 1
    if macd_status in ("金叉(转多)", "多头持续"):
        score += 1
    if rsi_flag == "中性":
        score += 1
    if above_ma20 == "是":
        score += 1
    if rel_strength is not None and rel_strength > 0:
        score += 1

    # ========== 量化打分引擎：趋势分(25) + 量能分(17) + 板块分(12) - 风险扣分 ==========
    # 说明见文件头部注释；这是常见多因子选股框架的简化落地版，不含大盘环境和竞价模块

    # ---- 趋势得分 TrendScore(满分25) ----
    ma10 = calc_ma(df, 10).iloc[-1]
    trend_score = 0
    t_detail = []

    if not pd.isna(ma10) and ma5 > ma10 > ma20:
        trend_score += 5; t_detail.append("均线多头排列(MA5>MA10>MA20)")
    if latest["收盘"] > ma20:
        trend_score += 3; t_detail.append("股价站上20日均线")
    ma20_series = calc_ma(df, 20)
    if len(ma20_series) > 25 and not pd.isna(ma20_series.iloc[-6]) and ma20 > ma20_series.iloc[-6]:
        trend_score += 3; t_detail.append("20日均线自身向上")
    if len(df) >= 21 and latest["收盘"] > df["最高"].iloc[-21:-1].max():
        trend_score += 4; t_detail.append("突破20日新高")
    if len(df) >= 61 and latest["收盘"] > df["最高"].iloc[-61:-1].max():
        trend_score += 3; t_detail.append("突破60日新高")
    if rel_strength is not None and rel_strength > 10:
        trend_score += 3; t_detail.append("20日跑赢沪深300超过10%")
    if bias20 is not None:
        if latest["收盘"] < ma20:
            pass
        elif bias20 <= 5:
            trend_score += 4; t_detail.append("距20日均线0-5%(不过度偏离)")
        elif bias20 <= 10:
            trend_score += 3; t_detail.append("距20日均线5-10%")
        elif bias20 <= 15:
            trend_score += 1; t_detail.append("距20日均线10-15%")

    # ---- 量能得分 VolumeScore(满分17，原框架含"主力资金流"3分因数据源限制未实现) ----
    volume_score = 0
    v_detail = []
    ma20_vol = df["成交量"].rolling(20).mean().iloc[-1] if "成交量" in df.columns else None
    today_vol_ratio = None
    if ma20_vol and ma20_vol > 0:
        today_vol_ratio = latest["成交量"] / ma20_vol
        if today_vol_ratio < 0.8:
            pass
        elif today_vol_ratio <= 1.2:
            volume_score += 2; v_detail.append("量能温和(0.8-1.2倍)")
        elif today_vol_ratio <= 1.5:
            volume_score += 4; v_detail.append("量能放大(1.2-1.5倍)")
        elif today_vol_ratio <= 2.5:
            volume_score += 6; v_detail.append("明显放量(1.5-2.5倍)")
        else:
            volume_score += 4; v_detail.append("极端放量(>2.5倍，需留意)")

    turnover_pct_rank = None
    if "换手率" in df.columns and len(df) >= 21:
        turnover_pct_rank = percentile_rank(df["换手率"].iloc[-21:-1], latest["换手率"])
        if turnover_pct_rank is not None and 60 <= turnover_pct_rank <= 90:
            volume_score += 3; v_detail.append("换手率处于近20日60-90分位(活跃但不过热)")

    if chg_5 is not None and len(df) >= 2:
        today_return = (latest["收盘"] - df["收盘"].iloc[-2]) / df["收盘"].iloc[-2] * 100
        ma5_vol = df["成交量"].rolling(5).mean().iloc[-1] if "成交量" in df.columns else None
        if today_return > 0 and ma5_vol and latest["成交量"] > ma5_vol:
            volume_score += 3; v_detail.append("今日上涨且放量")

    if "成交量" in df.columns and len(df) >= 6:
        last5 = df.iloc[-5:].copy()
        last5["前收"] = df["收盘"].shift(1).iloc[-5:]
        up_vol = last5[last5["收盘"] > last5["前收"]]["成交量"].mean()
        down_vol = last5[last5["收盘"] < last5["前收"]]["成交量"].mean()
        if pd.notna(up_vol) and pd.notna(down_vol) and up_vol > down_vol:
            volume_score += 3; v_detail.append("近5日上涨日成交量大于下跌日")

    if "成交量" in df.columns and len(df) >= 21:
        obv = calc_obv(df)
        if obv.iloc[-1] >= obv.iloc[-20:].max():
            volume_score += 2; v_detail.append("OBV创20日新高")

    # ---- 板块得分 SectorScore(简化版，满分12，只用板块自身涨跌幅，未做全市场板块横向排名) ----
    sector_score = 0
    s_detail = []
    sector_name = None
    sector_chg_1 = None
    sector_chg_5 = None
    sector_chg_20 = None

    if sector_info is not None:
        sector_name = sector_info.get("行业名称")

    if sector_hist is not None and len(sector_hist) >= 21:
        s_close = sector_hist["收盘"]
        sector_chg_1 = round((s_close.iloc[-1] - s_close.iloc[-2]) / s_close.iloc[-2] * 100, 2)
        sector_chg_5 = round((s_close.iloc[-1] - s_close.iloc[-6]) / s_close.iloc[-6] * 100, 2)
        sector_chg_20 = round((s_close.iloc[-1] - s_close.iloc[-21]) / s_close.iloc[-21] * 100, 2)

        if sector_chg_1 > 2:
            sector_score += 5; s_detail.append("板块当日涨幅>2%")
        elif sector_chg_1 > 0.5:
            sector_score += 3; s_detail.append("板块当日涨幅0.5%-2%")
        elif sector_chg_1 > -0.5:
            sector_score += 1; s_detail.append("板块当日表现平稳")

        if sector_chg_5 > 5:
            sector_score += 4; s_detail.append("板块5日涨幅>5%")
        elif sector_chg_5 > 1:
            sector_score += 2; s_detail.append("板块5日涨幅1%-5%")

        if chg_20 is not None and chg_20 > sector_chg_20:
            sector_score += 3; s_detail.append("个股20日涨幅跑赢所属板块")

    # ---- 风险扣分 RiskPenalty(实现R1/R2/R3，R4-R6因需要更多数据暂未实现) ----
    risk_penalty = 0
    r_detail = []
    if len(df) >= 4:
        chg_3 = (latest["收盘"] - df["收盘"].iloc[-4]) / df["收盘"].iloc[-4] * 100
        if chg_3 > 25:
            risk_penalty += 5; r_detail.append("最近3日累计涨幅超过25%")
    if bias20 is not None and bias20 > 20:
        risk_penalty += 5; r_detail.append("严重偏离20日均线(>20%)")
    if "换手率" in df.columns and len(df) >= 61:
        turnover_60_rank = percentile_rank(df["换手率"].iloc[-61:-1], latest["换手率"])
        if turnover_60_rank is not None and turnover_60_rank > 95:
            risk_penalty += 3; r_detail.append("换手率处于近60日95分位以上(极端活跃)")

    # ---- 汇总成 0-100 综合评分 ----
    raw_total = trend_score + volume_score + sector_score
    raw_max = 25 + 17 + 12  # =54，已实现模块的满分之和
    after_risk = max(0, raw_total - risk_penalty)
    final_score_100 = round(after_risk / raw_max * 100, 1)

    return {
        "股票代码": code,
        "股票名称": name,
        "最新收盘": latest["收盘"],
        "5日涨跌幅%": chg_5,
        "20日涨跌幅%": chg_20,
        "60日涨跌幅%": chg_60,
        "量能比(5日/前20日)": vol_ratio,
        "均线排列": ma_trend,
        "是否站上20日线": above_ma20,
        "乖离率%(距20日线)": bias20,
        "MACD状态": macd_status,
        "RSI(14)": rsi,
        "RSI区间": rsi_flag,
        "相对沪深300强弱%(20日)": rel_strength,
        "ATR止损参考幅度%": stop_loss_pct,
        "止损参考价": stop_loss_price,
        "20日均线值": round(ma20, 2) if ma20 and not pd.isna(ma20) else None,
        "60日最高": high_60,
        "60日最低": low_60,
        "综合趋势评分(0-5)": score,
        "最新日期": latest["日期"].strftime("%Y-%m-%d"),
        # ---- 量化打分引擎新增字段 ----
        "所属板块": sector_name,
        "板块当日涨跌幅%": sector_chg_1,
        "板块5日涨跌幅%": sector_chg_5,
        "趋势得分(0-25)": trend_score,
        "量能得分(0-17)": volume_score,
        "板块得分(0-12)": sector_score,
        "风险扣分": risk_penalty,
        "综合评分(0-100)": final_score_100,
        "评分明细_趋势": "；".join(t_detail) if t_detail else "无明显信号",
        "评分明细_量能": "；".join(v_detail) if v_detail else "无明显信号",
        "评分明细_板块": "；".join(s_detail) if s_detail else "无板块数据或无明显信号",
        "评分明细_风险": "；".join(r_detail) if r_detail else "无风险扣分项触发",
    }


def prepare_ranked(result_df: pd.DataFrame) -> pd.DataFrame:
    """按综合评分(0-100)+5日涨幅排序，并加上排名列，Excel和网页版JSON共用这份结果"""
    df = result_df.sort_values(
        ["综合评分(0-100)", "5日涨跌幅%"], ascending=[False, False], na_position="last"
    ).reset_index(drop=True)
    df.insert(0, "综合排名", range(1, len(df) + 1))
    return df


def export_json(result_df: pd.DataFrame, json_path: str):
    """导出给网页版仪表盘用的JSON数据文件"""
    import json
    from datetime import datetime, timezone, timedelta

    beijing_now = datetime.now(timezone(timedelta(hours=8)))
    records = json.loads(result_df.to_json(orient="records", force_ascii=False))
    latest_trade_dates = [r.get("最新日期") for r in records if r.get("最新日期")]
    data_trade_date = max(latest_trade_dates) if latest_trade_dates else None
    payload = {
        "生成时间": beijing_now.strftime("%Y-%m-%d %H:%M:%S") + "(北京时间)",
        "数据对应交易日": data_trade_date,
        "股票数量": len(records),
        "数据": records,
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def style_and_save(result_df: pd.DataFrame):
    result_df.to_excel(OUTPUT_FILE, index=False, sheet_name="选股分析")

    from openpyxl import load_workbook
    wb = load_workbook(OUTPUT_FILE)
    ws = wb["选股分析"]

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    header_font = Font(bold=True)

    for cell in ws[1]:
        cell.font = header_font

    col_index_map = {cell.value: cell.column for cell in ws[1]}

    # 数值型涨跌类列：正数红、负数绿
    numeric_color_cols = [
        "5日涨跌幅%", "20日涨跌幅%", "60日涨跌幅%",
        "乖离率%(距20日线)", "相对沪深300强弱%(20日)",
    ]
    for col_name in numeric_color_cols:
        if col_name not in col_index_map:
            continue
        col_letter = get_column_letter(col_index_map[col_name])
        for row in range(2, ws.max_row + 1):
            cell = ws[f"{col_letter}{row}"]
            if isinstance(cell.value, (int, float)):
                cell.fill = red_fill if cell.value > 0 else (green_fill if cell.value < 0 else cell.fill)

    # 文字状态列：多头/金叉标红，空头/死叉标绿
    text_color_cols = ["均线排列", "MACD状态"]
    bullish_words = ["多头", "金叉"]
    bearish_words = ["空头", "死叉"]
    for col_name in text_color_cols:
        if col_name not in col_index_map:
            continue
        col_letter = get_column_letter(col_index_map[col_name])
        for row in range(2, ws.max_row + 1):
            cell = ws[f"{col_letter}{row}"]
            val = str(cell.value)
            if any(w in val for w in bullish_words):
                cell.fill = red_fill
            elif any(w in val for w in bearish_words):
                cell.fill = green_fill

    # RSI区间：超买/超卖标黄提醒关注
    if "RSI区间" in col_index_map:
        col_letter = get_column_letter(col_index_map["RSI区间"])
        for row in range(2, ws.max_row + 1):
            cell = ws[f"{col_letter}{row}"]
            if cell.value in ("超买区间", "超卖区间"):
                cell.fill = yellow_fill

    # 综合评分列：4-5分标红，0-1分标绿
    if "综合趋势评分(0-5)" in col_index_map:
        col_letter = get_column_letter(col_index_map["综合趋势评分(0-5)"])
        for row in range(2, ws.max_row + 1):
            cell = ws[f"{col_letter}{row}"]
            if isinstance(cell.value, (int, float)):
                if cell.value >= 4:
                    cell.fill = red_fill
                elif cell.value <= 1:
                    cell.fill = green_fill

    # 新版综合评分(0-100)：>=70标红，<=30标绿
    if "综合评分(0-100)" in col_index_map:
        col_letter = get_column_letter(col_index_map["综合评分(0-100)"])
        for row in range(2, ws.max_row + 1):
            cell = ws[f"{col_letter}{row}"]
            if isinstance(cell.value, (int, float)):
                if cell.value >= 70:
                    cell.fill = red_fill
                elif cell.value <= 30:
                    cell.fill = green_fill

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    # 加一个说明sheet，把免责声明和指标含义写清楚
    ws2 = wb.create_sheet("指标说明")
    notes = [
        ["重要说明", "本报表所有指标基于公开技术分析公式计算，是历史价格的客观统计描述，"
                     "不构成投资建议，不保证未来表现。是否建仓、仓位大小、止损位设置，"
                     "最终需要你自己判断和承担风险。"],
        ["", ""],
        ["指标", "含义"],
        ["5/20/60日涨跌幅", "近N个交易日收盘价的涨跌百分比"],
        ["量能比", "近5日平均成交量 相比 之前20日平均成交量 的倍数，大于1说明近期在放量"],
        ["均线排列", "5/20/60日均线的排列方向，多头排列=短期在上长期在下且趋势向上"],
        ["是否站上20日线", "最新收盘价是否在20日均线上方"],
        ["乖离率", "最新收盘价偏离20日均线的百分比，绝对值越大代表离均线越远，可能有回归需求"],
        ["MACD状态", "基于DIF/DEA判断的金叉/死叉/多空持续状态"],
        ["RSI(14)", "14日相对强弱指标，>=70通常视为超买，<=30通常视为超卖"],
        ["相对沪深300强弱", "该股近20日涨跌幅 减去 沪深300同期涨跌幅，正数代表跑赢大盘"],
        ["ATR止损参考幅度", "基于近14日真实波动幅度(ATR)的2倍，计算出的止损参考百分比，"
                          "是常见的技术止损方法之一，不是必须遵循的规则"],
        ["综合趋势评分", "0-5分，统计多头排列/MACD偏多/RSI中性/站上20日线/跑赢大盘 这5项"
                       "各是否成立，分数越高代表当前技术面信号偏多，仅供参考"],
        ["", ""],
        ["== 量化打分引擎(0-100) ==", "参考常见的多因子选股框架，把趋势/量能/板块拆解成具体加分项。"
                                    "只实现了个股趋势、量能、板块相对表现三大模块和部分风险扣分项，"
                                    "不含大盘环境和集合竞价模块，分数换算到0-100仅为直观比较，"
                                    "不能和以后补全全部模块的分数直接对比。评分高不等于上涨概率高，"
                                    "两者是不同概念，仅供参考，不构成投资建议。"],
        ["趋势得分(0-25)", "均线多头排列/站上20日线/均线向上/突破20日或60日新高/跑赢大盘/"
                          "合理乖离幅度 等信号加总"],
        ["量能得分(0-17)", "成交量倍数/换手率合理分位/上涨放量/近5日量价关系/OBV创新高 等信号加总"],
        ["板块得分(0-12)", "所属申万行业板块当日及5日涨跌幅、个股是否跑赢所属板块"],
        ["风险扣分", "3日涨幅过大/严重偏离20日均线/换手率处于历史极端分位 等风险信号扣分"],
        ["综合评分(0-100)", "(趋势得分+量能得分+板块得分-风险扣分)/54 换算成的0-100分，"
                           "分数越高代表当前技术面信号越集中偏多，不是上涨概率，仅供参考"],
        ["评分明细_XX", "具体触发了哪些加分/扣分项，方便理解分数是怎么来的"],
    ]
    for row in notes:
        ws2.append(row)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 90
    for cell in ws2["A"]:
        cell.font = Font(bold=True)
    ws2["A1"].fill = yellow_fill
    ws2["B1"].fill = yellow_fill

    wb.save(OUTPUT_FILE)


def main():
    print("正在读取 history_data 里的数据...")
    all_files = load_all_files()

    benchmark_df = all_files.pop(BENCHMARK_CODE, None)
    if benchmark_df is None:
        print("提示：没有找到沪深300数据(000300_沪深300.csv)，相对强弱这项会缺失。")
        print("      可以重新运行 fetch_history.py 来补充抓取。")

    stock_to_sector, sector_histories = load_sector_data()
    if not stock_to_sector:
        print("提示：没有找到板块数据，板块得分这项会跳过（不影响其他指标）。")
        print("      可以运行一次 fetch_sector.py 来补充抓取。")

    if not all_files:
        print("没有找到任何个股历史数据，请先运行 fetch_history.py")
        return

    print(f"共读取到 {len(all_files)} 只股票，正在计算技术指标和量化评分...")
    rows = []
    for code, df in all_files.items():
        sector_info = stock_to_sector.get(code)
        sector_hist = None
        if sector_info:
            sector_hist = sector_histories.get(sector_info["行业代码"])
        metrics = calc_metrics(df, benchmark_df, sector_info, sector_hist)
        if metrics:
            rows.append(metrics)
        else:
            print(f"  {code} 数据不足30天，跳过（建议数据积累更长时间后再分析）")

    if not rows:
        print("数据不足，无法计算指标")
        return

    result_df = pd.DataFrame(rows)
    ranked_df = prepare_ranked(result_df)
    style_and_save(ranked_df)
    export_json(ranked_df, JSON_FILE)

    print(f"\n完成！报表已生成：{OUTPUT_FILE}")
    print(f"网页版数据已生成：{JSON_FILE}（打开 dashboard.html 后选择这个文件即可查看）")
    print("已按 综合评分(0-100) + 5日涨跌幅 排序")
    print("详细指标含义见 Excel 里的\"指标说明\"sheet")
    print("\n预览评分前5名：")
    cols = ["股票名称", "综合评分(0-100)", "趋势得分(0-25)", "量能得分(0-17)", "板块得分(0-12)", "风险扣分", "所属板块"]
    print(ranked_df.head(5)[cols])


if __name__ == "__main__":
    main()
